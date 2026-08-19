from fastapi import APIRouter, UploadFile, File, Form, HTTPException

from app.gemini import generate_response

from pptx import Presentation
from PyPDF2 import PdfReader

import io

router = APIRouter(tags=["Presentation"])


@router.post("/presentation/analyze")
async def analyze_presentation(
    file: UploadFile | None = File(None),
    transcript: str = Form("")
):
    try:
        extracted_text = ""

        # Read uploaded file
        if file:
            contents = await file.read()

            filename = (file.filename or "").lower()

            # PDF
            if filename.endswith(".pdf"):
                pdf = PdfReader(io.BytesIO(contents))

                for page in pdf.pages:
                    text = page.extract_text()

                    if text:
                        extracted_text += text + "\n"

            # PowerPoint
            elif filename.endswith(".pptx"):
                ppt = Presentation(io.BytesIO(contents))

                for slide in ppt.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            extracted_text += shape.text + "\n"

            # Text
            elif filename.endswith(".txt"):
                extracted_text = contents.decode(
                    "utf-8",
                    errors="ignore"
                )

            else:
                raise HTTPException(
                    status_code=400,
                    detail="Supported files: PDF, PPTX and TXT."
                )

        prompt = f"""
You are Oratio AI, a professional presentation and communication coach.

Analyze the available presentation content and speaker transcript.

PRESENTATION CONTENT:
{extracted_text[:16000]}

SPEAKER TRANSCRIPT:
{transcript[:16000]}

Use whichever inputs are available.

If only presentation content is available:
evaluate content, structure, clarity, and organization.

If only transcript is available:
evaluate speaking and communication performance.

If both are available:
evaluate both and compare the spoken content with the presentation content.

Return ONLY this plain-text format:

ORATIO AI PRESENTATION EVALUATION

OVERALL SCORE
__/100

CONFIDENCE SCORE
__/100

CLARITY SCORE
__/100

GRAMMAR SCORE
__/100

PRESENTATION STRUCTURE
__/100

SPEAKING PACE
__/100

AUDIENCE ENGAGEMENT
__/100

FILLER WORD USAGE
Low / Medium / High

CONTENT QUALITY
__/100

STRENGTHS

1.
2.
3.

AREAS FOR IMPROVEMENT

1.
2.
3.

RECOMMENDATIONS

1.
2.
3.
4.
5.

CONTENT COVERAGE

If both presentation content and transcript are available, explain whether the speaker covered the important presentation points.

OVERALL SUMMARY

Write one concise professional paragraph.

RULES:
- Always provide every section.
- Always provide numerical scores when meaningful.
- Do not use emojis.
- Do not use markdown bold.
- Do not use asterisks.
- Do not add an introduction.
- Do not ask for additional information.
"""

        feedback = generate_response(prompt)

        return {
            "feedback": feedback
        }

    except HTTPException:
        raise

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=str(e)
        )