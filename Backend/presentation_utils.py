from pptx import Presentation
import pdfplumber


def extract_ppt_text(file_path):
    """
    Extract all text from a PowerPoint presentation.
    """

    presentation = Presentation(file_path)

    text = ""

    for slide in presentation.slides:

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text += shape.text + "\n"

    return text


def extract_pdf_text(file_path):
    """
    Extract all text from a PDF.
    """

    text = ""

    with pdfplumber.open(file_path) as pdf:

        for page in pdf.pages:

            page_text = page.extract_text()

            if page_text:
                text += page_text + "\n"

    return text